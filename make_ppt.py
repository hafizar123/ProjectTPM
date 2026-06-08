from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

# Warna tema Bersih.In
TOSCA_DARK   = RGBColor(0x02, 0x59, 0x55)
TOSCA_MED    = RGBColor(0x00, 0x90, 0x9E)
TOSCA_LIGHT  = RGBColor(0x48, 0xC9, 0xB0)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT    = RGBColor(0x1A, 0x1A, 0x2E)
LIGHT_BG     = RGBColor(0xF0, 0xF9, 0xF8)
GRAY_TEXT    = RGBColor(0x55, 0x65, 0x7A)
ACCENT       = RGBColor(0xFF, 0xC1, 0x07)

BLANK = prs.slide_layouts[6]  # blank layout

def add_rect(slide, l, t, w, h, fill=None, line=None, line_w=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background() if line is None else None
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        if line_w:
            shape.line.width = line_w
    else:
        shape.line.fill.background()
    return shape

def add_text_box(slide, text, l, t, w, h, size=18, bold=False, color=None,
                 align=PP_ALIGN.LEFT, wrap=True, italic=False, spacing_after=0):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    if spacing_after:
        p.space_after = Pt(spacing_after)
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    run.font.color.rgb = color or DARK_TEXT
    return txb

def add_bullet_box(slide, lines, l, t, w, h, size=14, color=None, title=None,
                   title_size=15, title_color=None, bullet_char="•"):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = txb.text_frame
    tf.word_wrap = True
    first = True
    if title:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = title
        run.font.size  = Pt(title_size)
        run.font.bold  = True
        run.font.color.rgb = title_color or TOSCA_DARK
        p.space_after = Pt(4)
    for line in lines:
        p = tf.paragraphs[0] if (first and not title) else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = f"{bullet_char}  {line}"
        run.font.size  = Pt(size)
        run.font.color.rgb = color or DARK_TEXT
        p.space_after = Pt(3)
    return txb

def header_bar(slide, title, subtitle=None):
    add_rect(slide, 0, 0, 13.33, 1.35, fill=TOSCA_DARK)
    add_rect(slide, 0, 1.35, 13.33, 0.07, fill=TOSCA_LIGHT)
    add_text_box(slide, title, 0.5, 0.15, 12, 0.75,
                 size=32, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    if subtitle:
        add_text_box(slide, subtitle, 0.5, 0.82, 12, 0.45,
                     size=15, color=TOSCA_LIGHT, align=PP_ALIGN.LEFT, italic=True)

def footer_bar(slide, text="Bersih.In  •  Teknologi Pemrograman Mobile  •  2025/2026"):
    add_rect(slide, 0, 7.15, 13.33, 0.35, fill=TOSCA_DARK)
    add_text_box(slide, text, 0.3, 7.17, 12.5, 0.3,
                 size=10, color=TOSCA_LIGHT, align=PP_ALIGN.LEFT)

def card(slide, l, t, w, h, fill=LIGHT_BG, radius=None):
    sh = add_rect(slide, l, t, w, h, fill=fill, line=TOSCA_LIGHT, line_w=Pt(0.75))
    return sh

# ─────────────────────────────────────────────────────────────
# SLIDE 1 — COVER
# ─────────────────────────────────────────────────────────────
s1 = prs.slides.add_slide(BLANK)
add_rect(s1, 0, 0, 13.33, 7.5, fill=TOSCA_DARK)
add_rect(s1, 0, 0, 13.33, 0.08, fill=TOSCA_LIGHT)
add_rect(s1, 0, 7.42, 13.33, 0.08, fill=TOSCA_LIGHT)
# dekor lingkaran
for cx, cy, r in [(11.5,1.2,1.8),(12.2,5.5,2.2),(1.0,6.5,1.2)]:
    sh = s1.shapes.add_shape(9, Inches(cx-r/2), Inches(cy-r/2), Inches(r), Inches(r))
    sh.fill.solid(); sh.fill.fore_color.rgb = TOSCA_MED
    sh.line.fill.background()
    sh.fill.fore_color.theme_color  # just touch it
    # manual opacity trick via XML
    from lxml import etree
    sp_pr = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
    if sp_pr is not None:
        srgb = sp_pr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        if srgb is not None:
            alpha = etree.SubElement(srgb, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
            alpha.set('val', '15000')

add_text_box(s1, "Bersih.In", 1.5, 1.8, 10, 1.2,
             size=72, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text_box(s1, "Your Clean Space, Perfected", 1.5, 2.9, 10, 0.6,
             size=22, color=TOSCA_LIGHT, align=PP_ALIGN.CENTER, italic=True)

add_rect(s1, 3.8, 3.65, 5.7, 0.05, fill=TOSCA_LIGHT)

add_text_box(s1, "Aplikasi Layanan Kebersihan & Perawatan Hunian Berbasis Mobile",
             1.5, 3.85, 10, 0.55, size=16, color=WHITE, align=PP_ALIGN.CENTER)

add_text_box(s1,
             "Akmal Danendra Maulana  (123230135)   •   Hafiz Alaudin Rasendriya  (123230149)",
             1.5, 4.6, 10, 0.45, size=13, color=TOSCA_LIGHT, align=PP_ALIGN.CENTER)
add_text_box(s1, "Teknologi Pemrograman Mobile  •  2025 / 2026",
             1.5, 5.0, 10, 0.4, size=13, color=WHITE, align=PP_ALIGN.CENTER)

# ─────────────────────────────────────────────────────────────
# SLIDE 2 — LATAR BELAKANG & DESKRIPSI PROYEK
# ─────────────────────────────────────────────────────────────
s2 = prs.slides.add_slide(BLANK)
add_rect(s2, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s2, "Latar Belakang & Deskripsi Proyek",
           "Mengapa Bersih.In dibangun?")
footer_bar(s2)

# Masalah
card(s2, 0.4, 1.6, 5.8, 2.5)
add_text_box(s2, "❌  Masalah", 0.6, 1.7, 5.4, 0.45, size=16, bold=True, color=TOSCA_DARK)
add_bullet_box(s2,
    ["Sulit menemukan jasa kebersihan profesional yang terpercaya",
     "Tidak ada transparansi harga & status pengerjaan",
     "Tidak ada platform terintegrasi untuk berbagai layanan hunian",
     "Pengguna kesulitan menjadwal & memantau teknisi"],
    0.6, 2.15, 5.5, 1.9, size=13, color=DARK_TEXT)

# Solusi
card(s2, 6.8, 1.6, 5.9, 2.5)
add_text_box(s2, "✅  Solusi", 7.0, 1.7, 5.5, 0.45, size=16, bold=True, color=TOSCA_DARK)
add_bullet_box(s2,
    ["Platform mobile yang menghubungkan pengguna & teknisi terverifikasi",
     "Pemesanan mudah dengan pilihan jadwal & alamat GPS",
     "Transparansi harga & tracking status real-time",
     "Didukung AI assistant & rekomendasi layanan cerdas"],
    7.0, 2.15, 5.6, 1.9, size=13, color=DARK_TEXT)

# Deskripsi
card(s2, 0.4, 4.3, 12.3, 2.55)
add_text_box(s2, "Tentang Bersih.In", 0.65, 4.4, 11.8, 0.45,
             size=16, bold=True, color=TOSCA_DARK)
add_text_box(s2,
    "Bersih.In adalah aplikasi layanan kebersihan dan perawatan hunian berbasis mobile yang beroperasi di Yogyakarta. "
    "Platform ini menyediakan 8 kategori layanan — mulai dari Reguler Cleaning, Deep Cleaning, Service AC, Cuci Kasur, hingga Pijat Relaksasi — "
    "dengan sistem pemesanan terintegrasi, pembayaran multi-metode, dan manajemen karyawan otomatis.",
    0.65, 4.82, 11.8, 1.8, size=13, color=DARK_TEXT, wrap=True)

# ─────────────────────────────────────────────────────────────
# SLIDE 3 — TECH STACK
# ─────────────────────────────────────────────────────────────
s3 = prs.slides.add_slide(BLANK)
add_rect(s3, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s3, "Tech Stack", "Teknologi yang digunakan dalam pembangunan Bersih.In")
footer_bar(s3)

tech_cols = [
    ("📱  Frontend", TOSCA_DARK, [
        "Flutter (Dart)  —  UI cross-platform",
        "Hive  —  local key-value storage",
        "SharedPreferences  —  session & settings",
        "flutter_local_notifications  —  push notif lokal",
        "sensors_plus  —  gyroscope & akselerometer",
        "google_fonts / Google Generative AI SDK",
    ]),
    ("⚙️  Backend", TOSCA_MED, [
        "Node.js + Express.js  —  REST API server",
        "MySQL 8  —  relational database",
        "mysql2  —  driver koneksi DB",
        "bcryptjs  —  hash password",
        "jsonwebtoken  —  autentikasi JWT",
        "dotenv  —  environment config",
    ]),
    ("🔌  Layanan Eksternal", RGBColor(0x2E,0x86,0xAB), [
        "Google Gemini API  —  AI chatbot assistant",
        "open.er-api.com  —  live exchange rates",
        "OpenStreetMap + Nominatim  —  peta & geocoding",
        "flutter_map + latlong2  —  rendering peta",
        "local_auth  —  biometrik sidik jari/wajah",
        "image_picker  —  pilih foto profil",
    ]),
]

for i, (title, color, items) in enumerate(tech_cols):
    x = 0.4 + i * 4.3
    card(s3, x, 1.6, 4.1, 5.2)
    add_text_box(s3, title, x+0.15, 1.7, 3.8, 0.45,
                 size=14, bold=True, color=color)
    add_rect(s3, x+0.15, 2.15, 3.7, 0.05, fill=color)
    add_bullet_box(s3, items, x+0.15, 2.3, 3.8, 4.4,
                   size=12, color=DARK_TEXT, bullet_char="▸")

# ─────────────────────────────────────────────────────────────
# SLIDE 4 — ARSITEKTUR MVC
# ─────────────────────────────────────────────────────────────
s4 = prs.slides.add_slide(BLANK)
add_rect(s4, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s4, "Arsitektur MVC", "Model – View – Controller")
footer_bar(s4)

layers = [
    ("VIEW",       "lib/views/\nlib/widgets/",    "Tampilan UI Flutter. Menerima input pengguna,\nmenampilkan data, tidak mengandung logika bisnis.",    TOSCA_LIGHT,  RGBColor(0x02,0x59,0x55)),
    ("CONTROLLER", "lib/controllers/",            "Jembatan View–Model. Memproses input, memanggil\nService, dan memperbarui state yang ditampilkan View.", TOSCA_MED,    WHITE),
    ("MODEL",      "lib/models/",                 "Representasi data: UserModel, OrderModel,\nAddressModel. Bertugas mapping JSON dari API.",            TOSCA_DARK,   WHITE),
    ("SERVICE",    "lib/services/",               "Komunikasi ke backend (HTTP), biometrik, notifikasi,\nAI, rekomendasi, dan konversi mata uang.",         RGBColor(0x1A,0x1A,0x2E), WHITE),
    ("BACKEND API","Node.js + MySQL",             "REST API Express.js. Validasi, hash password,\nquery DB, auto-assign karyawan, live exchange rates.",  RGBColor(0x0D,0x47,0xA1), WHITE),
]

box_w, box_h, gap = 2.1, 0.82, 0.1
start_x = 0.5
y_base  = 1.65

for i, (lbl, sub, desc, bg, fg) in enumerate(layers):
    x = start_x + i * (box_w + gap + 0.12)
    sh = add_rect(s4, x, y_base, box_w, box_h, fill=bg)
    add_text_box(s4, lbl, x, y_base+0.08, box_w, 0.36,
                 size=14, bold=True, color=fg, align=PP_ALIGN.CENTER)
    add_text_box(s4, sub, x, y_base+0.42, box_w, 0.35,
                 size=9, color=fg, align=PP_ALIGN.CENTER, italic=True)
    # arrow
    if i < len(layers)-1:
        ax = x + box_w + 0.03
        add_text_box(s4, "→", ax, y_base+0.28, 0.18, 0.3,
                     size=16, bold=True, color=TOSCA_DARK, align=PP_ALIGN.CENTER)

# deskripsi tiap layer di bawah
for i, (lbl, sub, desc, bg, fg) in enumerate(layers):
    x = start_x + i * (box_w + gap + 0.12)
    card(s4, x, y_base+0.98, box_w, 3.8, fill=WHITE)
    add_text_box(s4, desc, x+0.1, y_base+1.08, box_w-0.15, 3.6,
                 size=11, color=DARK_TEXT, wrap=True)

# folder tree
add_text_box(s4, "Struktur Folder Flutter", 0.5, 5.98, 5.5, 0.35,
             size=13, bold=True, color=TOSCA_DARK)
tree = ("lib/\n"
        "  controllers/   ← AuthController\n"
        "  models/        ← UserModel, OrderModel, AddressModel\n"
        "  services/      ← AuthService, AiService, CurrencyService…\n"
        "  views/         ← auth/, home/, order/, profile/, support/…\n"
        "  widgets/       ← CustomNavBar, HomeCarousel, SmartRecommendation…\n"
        "  main.dart")
add_text_box(s4, tree, 0.5, 6.3, 7.5, 1.15, size=10, color=DARK_TEXT,
             wrap=True)

# ─────────────────────────────────────────────────────────────
# SLIDE 5 — DATABASE
# ─────────────────────────────────────────────────────────────
s5 = prs.slides.add_slide(BLANK)
add_rect(s5, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s5, "Struktur Database", "MySQL — 9 tabel utama")
footer_bar(s5)

tables = [
    ("profiles",        "id, email, username, password, avatar_url"),
    ("employees",       "id, nama, username, email, password, no_hp, is_active"),
    ("orders",          "id, user_email, service_name, total_amount, currency,\ntotal_converted, payment_method, va_number, qris_url,\naddress, schedule_date, schedule_time, employee_id, status"),
    ("saved_addresses", "id, username, address, lat, lng, house_type, description"),
    ("messages",        "id, user_email, sender, message, created_at"),
    ("employee_chats",  "id, order_id, sender_role, sender_id, message, created_at"),
    ("order_reviews",   "id, order_id, user_email, rating, review, created_at"),
    ("evaluations",     "id, email, rating, kesan, saran, created_at"),
    ("reports",         "id, order_id, user_email, description, image_base64, status"),
]

cols = 3
rows_per_col = 3
for idx, (tbl, fields) in enumerate(tables):
    col = idx % cols
    row = idx // cols
    x = 0.4 + col * 4.3
    y = 1.65 + row * 1.85
    card(s5, x, y, 4.1, 1.75, fill=WHITE)
    add_text_box(s5, tbl, x+0.12, y+0.1, 3.85, 0.38,
                 size=13, bold=True, color=TOSCA_DARK)
    add_rect(s5, x+0.12, y+0.47, 3.6, 0.04, fill=TOSCA_LIGHT)
    add_text_box(s5, fields, x+0.12, y+0.55, 3.86, 1.1,
                 size=9.5, color=GRAY_TEXT, wrap=True)

# relasi
add_text_box(s5, "Relasi Utama:", 0.4, 7.0, 3, 0.3,
             size=11, bold=True, color=TOSCA_DARK)
add_text_box(s5,
    "orders.user_email → profiles.email   |   orders.employee_id → employees.id   |   "
    "order_reviews.order_id → orders.id   |   reports.order_id → orders.id   |   "
    "employee_chats.order_id → orders.id",
    0.4, 7.25, 12.5, 0.3, size=9.5, color=GRAY_TEXT, wrap=False)

# ─────────────────────────────────────────────────────────────
# SLIDE 6 — FITUR AUTENTIKASI
# ─────────────────────────────────────────────────────────────
s6 = prs.slides.add_slide(BLANK)
add_rect(s6, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s6, "Fitur Autentikasi", "Login, Register, Biometrik, Multi-Role")
footer_bar(s6)

auth_features = [
    ("📝  Register", [
        "Input email, username, password",
        "Password di-hash dengan bcryptjs",
        "Validasi di server sebelum simpan",
    ]),
    ("🔑  Login", [
        "Support email ATAU username",
        "JWT token, berlaku 1 jam",
        "Role detection: user / karyawan / admin",
    ]),
    ("🪪  Login Biometrik", [
        "Sidik jari atau pengenalan wajah",
        "Isolasi per-akun (multi-akun support)",
        "Dialog pilih akun jika >1 terdaftar",
    ]),
    ("👤  Multi-Role", [
        "User: pesan & pantau layanan",
        "Karyawan: terima & update order",
        "Admin: dashboard & manajemen penuh",
    ]),
]

for i, (title, items) in enumerate(auth_features):
    x = 0.4 + i * 3.15
    card(s6, x, 1.6, 2.9, 4.6)
    add_text_box(s6, title, x+0.12, 1.7, 2.65, 0.45,
                 size=14, bold=True, color=TOSCA_DARK)
    add_rect(s6, x+0.12, 2.13, 2.5, 0.05, fill=TOSCA_LIGHT)
    add_bullet_box(s6, items, x+0.12, 2.25, 2.65, 3.8,
                   size=12.5, color=DARK_TEXT)

# flow
add_text_box(s6, "Alur Autentikasi:", 0.4, 6.3, 3, 0.3,
             size=12, bold=True, color=TOSCA_DARK)
add_text_box(s6,
    "Buka App  →  Cek Session  →  Login / Biometrik  →  Role Check  →  Dashboard sesuai role",
    0.4, 6.6, 12.5, 0.4, size=13, color=TOSCA_MED, bold=True,
    align=PP_ALIGN.LEFT)

# ─────────────────────────────────────────────────────────────
# SLIDE 7 — PEMESANAN LAYANAN
# ─────────────────────────────────────────────────────────────
s7 = prs.slides.add_slide(BLANK)
add_rect(s7, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s7, "Fitur Pemesanan Layanan", "8 Layanan • GPS Address • Penjadwalan")
footer_bar(s7)

services = [
    ("🔥", "Pemanas Air",      "Rp 100k – 250k"),
    ("🧹", "Reguler Cleaning", "Rp 80k – 280k"),
    ("🚗", "Cuci Kendaraan",   "Rp 25k – 120k"),
    ("🛏️","Cuci Kasur",       "Rp 150k – 300k"),
    ("🏠", "Deep Cleaning",    "Rp 350k – 950k"),
    ("💆", "Pijat Relaksasi",  "Rp 100k – 200k"),
    ("❄️","Service AC",       "Rp 100k – 400k"),
    ("🛋️","Cuci Sofa",        "Rp 120k – 350k"),
]

for i, (ico, name, price) in enumerate(services):
    col = i % 4
    row = i // 4
    x = 0.4 + col * 3.1
    y = 1.65 + row * 1.4
    sh = add_rect(s7, x, y, 2.8, 1.25, fill=WHITE, line=TOSCA_LIGHT, line_w=Pt(0.75))
    add_text_box(s7, ico,   x+0.12, y+0.1, 0.5,  0.45, size=22, align=PP_ALIGN.LEFT)
    add_text_box(s7, name,  x+0.65, y+0.1, 2.05, 0.45, size=12, bold=True, color=TOSCA_DARK)
    add_text_box(s7, price, x+0.65, y+0.5, 2.05, 0.35, size=11, color=GRAY_TEXT)

# alur pesan
add_text_box(s7, "Alur Pemesanan", 0.4, 4.55, 4, 0.35,
             size=14, bold=True, color=TOSCA_DARK)
steps = [
    "① Pilih Layanan",
    "② Login",
    "③ Pilih / Tambah Alamat (GPS)",
    "④ Pilih Paket & Jadwal",
    "⑤ Pilih Pembayaran",
    "⑥ Konfirmasi Bayar",
]
for i, step in enumerate(steps):
    x = 0.4 + i * 2.1
    add_rect(s7, x, 5.0, 1.95, 0.7, fill=TOSCA_DARK if i % 2 == 0 else TOSCA_MED)
    add_text_box(s7, step, x+0.05, 5.08, 1.85, 0.55,
                 size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    if i < len(steps)-1:
        add_text_box(s7, "→", x+1.95, 5.17, 0.15, 0.35,
                     size=12, bold=True, color=TOSCA_DARK, align=PP_ALIGN.CENTER)

add_text_box(s7,
    "Catatan: Jadwal dipilih min. 1 jam ke depan  •  Jam operasional 07:00–21:00 WIB  •  Teknisi datang sesuai jadwal",
    0.4, 5.85, 12.5, 0.35, size=11, color=GRAY_TEXT, italic=True)

# ─────────────────────────────────────────────────────────────
# SLIDE 8 — PEMBAYARAN
# ─────────────────────────────────────────────────────────────
s8 = prs.slides.add_slide(BLANK)
add_rect(s8, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s8, "Fitur Pembayaran", "Multi-Metode • Kurs Live • Auto-Cancel")
footer_bar(s8)

pay_cols = [
    ("💳  QRIS", TOSCA_DARK, [
        "GoPay, Dana, OVO, ShopeePay",
        "Tampil QR Code di aplikasi",
        "Konfirmasi manual oleh admin",
    ]),
    ("🏦  Virtual Account", TOSCA_MED, [
        "BCA, Mandiri, BNI, BRI",
        "Nomor VA unik per transaksi",
        "Instruksi transfer ditampilkan",
    ]),
    ("🌍  Bank Internasional", RGBColor(0x2E,0x86,0xAB), [
        "Bank of China  (CNY — Yuan)",
        "United Overseas Bank  (SGD)",
        "Saudi National Bank  (SAR — Riyal)",
        "Kurs live dari open.er-api.com",
        "Cache 1 jam, fallback hardcoded",
    ]),
    ("⏱️  Auto-Cancel", RGBColor(0xD3,0x2F,0x2F), [
        "Countdown 30 menit",
        "Notifikasi peringatan di menit ke-20",
        "Order dibatalkan otomatis jika kadaluarsa",
        "Endpoint /api/orders/auto-cancel",
    ]),
]

for i, (title, color, items) in enumerate(pay_cols):
    x = 0.4 + i * 3.15
    card(s8, x, 1.6, 2.9, 5.15)
    add_text_box(s8, title, x+0.12, 1.7, 2.65, 0.45,
                 size=13, bold=True, color=color)
    add_rect(s8, x+0.12, 2.13, 2.5, 0.05, fill=color)
    add_bullet_box(s8, items, x+0.12, 2.25, 2.65, 4.4,
                   size=12, color=DARK_TEXT)

add_text_box(s8,
    "Konversi mata uang dilakukan real-time di sisi Flutter dengan cache 1 jam — tidak perlu API key di client",
    0.4, 6.9, 12.5, 0.35, size=11.5, color=TOSCA_DARK, italic=True,
    align=PP_ALIGN.LEFT)

# ─────────────────────────────────────────────────────────────
# SLIDE 9 — AI & REKOMENDASI
# ─────────────────────────────────────────────────────────────
s9 = prs.slides.add_slide(BLANK)
add_rect(s9, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s9, "Fitur AI & Rekomendasi Cerdas",
           "Google Gemini AI  •  Content-Based Filtering")
footer_bar(s9)

# AI Chat
card(s9, 0.4, 1.6, 6.0, 5.1)
add_text_box(s9, "🤖  AI Assistant (Gemini)", 0.6, 1.72, 5.7, 0.45,
             size=16, bold=True, color=TOSCA_DARK)
add_rect(s9, 0.6, 2.15, 5.5, 0.05, fill=TOSCA_LIGHT)
add_bullet_box(s9,
    ["Powered by Google Gemini Flash",
     "API key diambil dari backend — tidak hardcode di client",
     "System context berisi info lengkap: layanan, harga, cara pesan, kebijakan",
     "Menjawab pertanyaan seputar layanan dalam Bahasa Indonesia",
     "Dipanggil melalui floating action button di seluruh halaman",
     "Fallback message jika API tidak tersedia",
    ],
    0.6, 2.25, 5.7, 4.3, size=13, color=DARK_TEXT)

# CBF
card(s9, 6.8, 1.6, 6.1, 5.1)
add_text_box(s9, "📊  Content-Based Filtering", 7.0, 1.72, 5.8, 0.45,
             size=16, bold=True, color=TOSCA_DARK)
add_rect(s9, 7.0, 2.15, 5.5, 0.05, fill=TOSCA_LIGHT)
add_bullet_box(s9,
    ["Setiap layanan direpresentasikan sebagai vektor 8 dimensi:",
     "   [kebersihan, teknis, relaksasi, kendaraan, kasur/sofa,",
     "    harga rendah, harga sedang, harga tinggi]",
     "Histori order user diubah menjadi user profile vector",
     "Cosine Similarity menghitung kecocokan user ↔ layanan",
     "Top-3 layanan yang belum pernah dipesan direkomendasikan",
     "Guest: rekomendasi populer berdasarkan rating tertinggi",
     "Diimplementasi di Flutter (client) DAN di backend Node.js",
    ],
    7.0, 2.25, 5.85, 4.3, size=12.5, color=DARK_TEXT)

# ─────────────────────────────────────────────────────────────
# SLIDE 10 — MANAJEMEN PESANAN
# ─────────────────────────────────────────────────────────────
s10 = prs.slides.add_slide(BLANK)
add_rect(s10, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s10, "Manajemen Pesanan", "Tracking • Review • Laporan • Notifikasi")
footer_bar(s10)

# Status flow
add_text_box(s10, "Alur Status Pesanan:", 0.4, 1.62, 4, 0.35,
             size=13, bold=True, color=TOSCA_DARK)
statuses = [
    ("Menunggu\nPembayaran", TOSCA_DARK),
    ("Menunggu\nKonfirmasi",  TOSCA_MED),
    ("Sedang\nDikerjakan",    RGBColor(0x1A,0x78,0xC2)),
    ("Selesai",               RGBColor(0x2E,0x7D,0x32)),
    ("Dibatalkan",            RGBColor(0xC6,0x28,0x28)),
]
for i, (lbl, col) in enumerate(statuses):
    x = 0.4 + i * 2.45
    add_rect(s10, x, 2.05, 2.2, 0.75, fill=col)
    add_text_box(s10, lbl, x, 2.1, 2.2, 0.65,
                 size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    if i < len(statuses)-1:
        add_text_box(s10, "→", x+2.2, 2.28, 0.25, 0.3,
                     size=13, bold=True, color=TOSCA_DARK, align=PP_ALIGN.CENTER)

feat_cols = [
    ("📋  Riwayat Transaksi", [
        "Daftar semua pesanan user",
        "Filter by status",
        "Detail pesanan lengkap + nama karyawan",
        "Tombol review & laporan per order",
    ]),
    ("⭐  Review & Rating", [
        "Rating bintang 1–5",
        "Komentar teks bebas",
        "Satu review per order (unique constraint)",
        "Tampil di halaman publik About Us",
    ]),
    ("🚨  Laporan / Aduan", [
        "Deskripsi masalah teks",
        "Upload foto bukti (Base64)",
        "Status: pending → diproses → selesai",
        "Admin bisa update status laporan",
    ]),
    ("🔔  Notifikasi", [
        "Peringatan bayar (menit ke-20)",
        "Pembayaran kadaluarsa",
        "Pesanan dikonfirmasi / dikerjakan / selesai",
        "Pengingat H-1 dan hari-H jadwal",
        "Notifikasi promo untuk guest",
    ]),
]

for i, (title, items) in enumerate(feat_cols):
    x = 0.4 + i * 3.15
    card(s10, x, 3.0, 2.9, 3.7)
    add_text_box(s10, title, x+0.12, 3.1, 2.65, 0.45,
                 size=13, bold=True, color=TOSCA_DARK)
    add_rect(s10, x+0.12, 3.53, 2.5, 0.04, fill=TOSCA_LIGHT)
    add_bullet_box(s10, items, x+0.12, 3.62, 2.65, 3.0,
                   size=12, color=DARK_TEXT)

# ─────────────────────────────────────────────────────────────
# SLIDE 11 — LIVE CHAT & SUPPORT
# ─────────────────────────────────────────────────────────────
s11 = prs.slides.add_slide(BLANK)
add_rect(s11, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s11, "Live Chat & Support", "Chat User-Admin  •  Chat Per Order  •  Notifikasi")
footer_bar(s11)

chat_cards = [
    ("💬  Live Chat User-Admin",
     ["Percakapan real-time antara user & admin",
      "Admin melihat daftar semua room percakapan",
      "Unread count per room chat",
      "Tersimpan di tabel messages (MySQL)",
      "Tampil urut dari pesan terlama"]),
    ("🔧  Chat User-Karyawan",
     ["Chat khusus per order (tabel employee_chats)",
      "Sender role: 'user' atau 'employee'",
      "Read-only setelah order selesai/dibatalkan",
      "Karyawan bisa kirim update progress",
      "User bisa konfirmasi detail pekerjaan"]),
    ("🔔  Halaman Notifikasi",
     ["Riwayat semua notifikasi tersimpan lokal",
      "Dikelompokkan: Pembayaran, Pesanan, Jadwal",
      "Hapus satu / hapus semua",
      "Badge count di ikon notifikasi",
      "Diperbarui real-time saat notif masuk"]),
    ("🤖  AI Chat Assistant",
     ["Bisa diakses dari FAB di semua halaman",
      "Konteks sistem berisi daftar harga & layanan",
      "Quick suggestion pertanyaan umum",
      "Animasi typing indicator",
      "Riwayat chat dalam satu sesi"]),
]

for i, (title, items) in enumerate(chat_cards):
    x = 0.4 + i * 3.15
    card(s11, x, 1.6, 2.9, 5.1)
    add_text_box(s11, title, x+0.12, 1.7, 2.65, 0.45,
                 size=13, bold=True, color=TOSCA_DARK)
    add_rect(s11, x+0.12, 2.13, 2.5, 0.05, fill=TOSCA_LIGHT)
    add_bullet_box(s11, items, x+0.12, 2.25, 2.65, 4.3,
                   size=12.5, color=DARK_TEXT)

# ─────────────────────────────────────────────────────────────
# SLIDE 12 — ADMIN & KARYAWAN
# ─────────────────────────────────────────────────────────────
s12 = prs.slides.add_slide(BLANK)
add_rect(s12, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s12, "Fitur Admin & Karyawan",
           "Dashboard Admin  •  Auto-Assign  •  Dashboard Karyawan")
footer_bar(s12)

card(s12, 0.4, 1.6, 6.1, 5.1)
add_text_box(s12, "🖥️  Dashboard Admin", 0.6, 1.72, 5.8, 0.45,
             size=16, bold=True, color=TOSCA_DARK)
add_rect(s12, 0.6, 2.15, 5.5, 0.05, fill=TOSCA_LIGHT)
add_bullet_box(s12,
    ["Lihat semua pesanan dengan filter & detail lengkap",
     "Konfirmasi pembayaran → ubah status order",
     "Total pendapatan & jumlah pesanan selesai (revenue)",
     "Kelola semua laporan/aduan user + update status",
     "Balas pesan live chat dari semua user",
     "Login khusus admin (username: admin)",
     "Lihat semua evaluasi user",
    ],
    0.6, 2.25, 5.85, 4.3, size=13, color=DARK_TEXT)

card(s12, 6.8, 1.6, 6.1, 5.1)
add_text_box(s12, "👷  Dashboard Karyawan", 7.0, 1.72, 5.8, 0.45,
             size=16, bold=True, color=TOSCA_DARK)
add_rect(s12, 7.0, 2.15, 5.5, 0.05, fill=TOSCA_LIGHT)
add_bullet_box(s12,
    ["Login karyawan via email/password dari tabel employees",
      "Melihat daftar order yang di-assign kepadanya",
      "Update status: Sedang Dikerjakan → Selesai",
      "Chat langsung dengan user per order",
      "Auto-assign: sistem pilih karyawan beban paling ringan",
      "   • Algoritma: hitung load per slot waktu",
      "   • Jika beban sama → random terpilih",
      "   • Kapasitas max 10 order per slot waktu",
    ],
    7.0, 2.25, 5.85, 4.3, size=13, color=DARK_TEXT)

# ─────────────────────────────────────────────────────────────
# SLIDE 13 — FITUR TAMBAHAN
# ─────────────────────────────────────────────────────────────
s13 = prs.slides.add_slide(BLANK)
add_rect(s13, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
header_bar(s13, "Fitur Tambahan",
           "Jam Real-Time  •  Mini Game  •  Profil  •  Evaluasi TPM")
footer_bar(s13)

extra_cards = [
    ("🕐  Jam Real-Time Multi-Zona",
     ["Jam berjalan setiap detik di beranda",
      "4 zona: WIB, WITA, WIT, London (BST)",
      "Dropdown zona waktu, tersimpan di SharedPreferences",
      "Sapaan otomatis (Pagi/Siang/Sore/Malam)",
      "berdasarkan jam zona yang dipilih"]),
    ("🎮  Mini Game Tersembunyi",
     ["Guncang HP di beranda untuk membuka game",
      "Gyroscope untuk gerakkan sapu",
      "5 level, kotoran makin banyak tiap level",
      "Sistem skor +10 per kotoran",
      "Overlay timbul saat shake terdeteksi"]),
    ("👤  Manajemen Profil",
     ["Edit username, email, password",
      "Upload foto profil (kompresi + Base64)",
      "Cache foto per-akun (tetap ada post-logout)",
      "Tampilan avatar di AppBar & profil",
      "Hapus akun permanen"]),
    ("📝  Evaluasi TPM",
     ["Rating bintang 0.5–5.0 (half-star)",
      "Input kesan & saran teks bebas",
      "Ditampilkan di carousel 'Apa Kata Orang'",
      "Hanya untuk user yang sudah login",
      "Submit sekali, bisa isi ulang"]),
]

for i, (title, items) in enumerate(extra_cards):
    x = 0.4 + i * 3.15
    card(s13, x, 1.6, 2.9, 5.0)
    add_text_box(s13, title, x+0.12, 1.7, 2.65, 0.45,
                 size=13, bold=True, color=TOSCA_DARK)
    add_rect(s13, x+0.12, 2.13, 2.5, 0.05, fill=TOSCA_LIGHT)
    add_bullet_box(s13, items, x+0.12, 2.25, 2.65, 4.25,
                   size=12.5, color=DARK_TEXT)

# ─────────────────────────────────────────────────────────────
# SLIDE 14 — PENUTUP
# ─────────────────────────────────────────────────────────────
s14 = prs.slides.add_slide(BLANK)
add_rect(s14, 0, 0, 13.33, 7.5, fill=TOSCA_DARK)
add_rect(s14, 0, 0, 13.33, 0.08, fill=TOSCA_LIGHT)
add_rect(s14, 0, 7.42, 13.33, 0.08, fill=TOSCA_LIGHT)

for cx, cy, r in [(1.5,1.8,1.5),(12.0,6.0,2.0),(11.5,1.0,1.0)]:
    sh = s14.shapes.add_shape(9, Inches(cx-r/2), Inches(cy-r/2), Inches(r), Inches(r))
    sh.fill.solid(); sh.fill.fore_color.rgb = TOSCA_MED
    sh.line.fill.background()
    from lxml import etree
    sp_pr = sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
    if sp_pr is not None:
        srgb = sp_pr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        if srgb is not None:
            alpha = etree.SubElement(srgb, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
            alpha.set('val', '12000')

add_text_box(s14, "Kesimpulan", 1.5, 1.3, 10, 0.65,
             size=36, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_rect(s14, 4.0, 1.95, 5.3, 0.06, fill=TOSCA_LIGHT)

summary = [
    "✅   Aplikasi mobile lengkap untuk layanan kebersihan hunian",
    "✅   Arsitektur MVC yang bersih: Flutter + Node.js + MySQL",
    "✅   Fitur AI (Gemini) + Content-Based Filtering rekomendasi layanan",
    "✅   Multi-role: User, Karyawan, Admin dengan alur kerja terpisah",
    "✅   Pembayaran multi-metode termasuk valuta asing dengan kurs live",
    "✅   Notifikasi real-time, live chat, laporan, dan review terintegrasi",
    "✅   Fitur unik: jam multi-zona, mini game tersembunyi, biometrik multi-akun",
]
for i, line in enumerate(summary):
    add_text_box(s14, line, 1.5, 2.2 + i*0.55, 10.3, 0.48,
                 size=14, color=WHITE, align=PP_ALIGN.LEFT)

add_rect(s14, 4.0, 6.18, 5.3, 0.06, fill=TOSCA_LIGHT)
add_text_box(s14, "Terima Kasih", 1.5, 6.3, 10, 0.6,
             size=28, bold=True, color=TOSCA_LIGHT, align=PP_ALIGN.CENTER)

# ─────────────────────────────────────────────────────────────
OUT = r"d:\MOBILE\PROJEK\BersihIn_Presentasi.pptx"
prs.save(OUT)
print(f"Saved: {OUT}")
